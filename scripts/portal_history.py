"""Excelに保存されたアップロード実行記録から、直近1回の案内を作る。"""

import json
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


EXCEL_PATH = Path(
    r"C:\Users\susuk\OneDrive\勝北VBC\動画ポータル\動画管理.xlsx"
)
LOG_SHEET = "_portal_upload_log"
LOG_HEADERS = (
    "batchId", "batchStartedAt", "youtubeId", "confirmedAt",
    "orderInBatch", "eventDate", "eventName",
)
JST = timezone(timedelta(hours=9))


def now_jst():
    return datetime.now(JST).isoformat(timespec="microseconds")


def parse_timestamp(value):
    if not isinstance(value, str):
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            return None
        return parsed.timestamp()
    except (ValueError, OverflowError, OSError):
        return None


def normalize_event_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for pattern in ("%Y/%m/%d", "%Y-%m-%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def read_workbook_state(wb):
    ws = wb.active
    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[2] if cell.value is not None
    }
    if "youtube ID" not in headers or "no." not in headers:
        raise ValueError("管理表のアクティブシートに必要な見出しがありません。")
    saved_ids = set()
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[headers["no."] - 1] is None:
            continue
        video_id = str(row[headers["youtube ID"] - 1] or "").strip()
        if video_id:
            if video_id in saved_ids:
                raise ValueError(f"ExcelのYouTube IDが重複しています: {video_id}")
            saved_ids.add(video_id)

    records = []
    if LOG_SHEET not in wb.sheetnames:
        return saved_ids, records
    log = wb[LOG_SHEET]
    if log.max_column != len(LOG_HEADERS) or tuple(
        cell.value for cell in log[1]
    ) != LOG_HEADERS:
        raise ValueError(f"管理シート {LOG_SHEET} の形式が想定と異なります。")
    seen_ids = set()
    seen_orders = set()
    batch_starts = {}
    for row in log.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue
        record = dict(zip(LOG_HEADERS, row))
        video_id = record["youtubeId"]
        batch_id = record["batchId"]
        order = record["orderInBatch"]
        if (
            not isinstance(video_id, str) or video_id not in saved_ids
            or video_id in seen_ids or not isinstance(batch_id, str) or not batch_id
            or not isinstance(order, int) or isinstance(order, bool) or order < 1
            or (batch_id, order) in seen_orders
            or parse_timestamp(record["batchStartedAt"]) is None
            or parse_timestamp(record["confirmedAt"]) is None
            or not isinstance(record["eventName"], str) or not record["eventName"].strip()
        ):
            raise ValueError(f"アップロード記録が不正です: {video_id}")
        if record["eventDate"] is not None and (
            normalize_event_date(record["eventDate"]) != record["eventDate"]
        ):
            raise ValueError(f"アップロード記録の開催日が不正です: {video_id}")
        started_at = record["batchStartedAt"]
        if batch_id in batch_starts and batch_starts[batch_id] != started_at:
            raise ValueError(f"同じ実行IDの開始日時が一致しません: {batch_id}")
        batch_starts[batch_id] = started_at
        seen_ids.add(video_id)
        seen_orders.add((batch_id, order))
        records.append(record)
    return saved_ids, records


def read_excel_state(path=EXCEL_PATH):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return read_workbook_state(wb)
    finally:
        wb.close()


def append_upload_record(wb, *, batch_id, started_at, video_id,
                         confirmed_at, order, event_date, event_name):
    if LOG_SHEET not in wb.sheetnames:
        log = wb.create_sheet(LOG_SHEET)
        log.append(LOG_HEADERS)
        log.sheet_state = "hidden"
    log = wb[LOG_SHEET]
    log.append((
        batch_id, started_at, video_id, confirmed_at, order,
        normalize_event_date(event_date), str(event_name).strip(),
    ))


def build_latest_news(records, visible_ids):
    batches = {}
    for record in records:
        batches.setdefault(record["batchId"], []).append(record)
    candidates = []
    for batch_id, batch in batches.items():
        missing = {item["youtubeId"] for item in batch} - visible_ids
        if missing:
            print(f"新着情報を保留: 実行 {batch_id} の未掲載動画 {len(missing)}本")
            continue
        latest = max(batch, key=lambda item: parse_timestamp(item["confirmedAt"]))
        events = {}
        for item in sorted(batch, key=lambda record: record["orderInBatch"]):
            key = (item["eventDate"], item["eventName"])
            if key not in events:
                events[key] = {
                    "eventDate": item["eventDate"],
                    "eventName": item["eventName"],
                    "videoCount": 0,
                }
            events[key]["videoCount"] += 1
        candidates.append({
            "batchId": batch_id,
            "uploadedAt": latest["confirmedAt"],
            "videoCount": len(batch),
            "events": list(events.values()),
        })
    candidates.sort(
        key=lambda item: (parse_timestamp(item["uploadedAt"]), item["batchId"]),
        reverse=True,
    )
    return candidates[:1]


def has_pending_portal_update(wb, output_file):
    saved_ids, records = read_workbook_state(wb)
    if not output_file.exists():
        return bool(saved_ids)
    output = json.loads(output_file.read_text(encoding="utf-8"))
    visible_ids = {video["id"] for video in output["videos"]}
    return bool(saved_ids - visible_ids) or (
        output.get("news", []) != build_latest_news(records, visible_ids)
    )
