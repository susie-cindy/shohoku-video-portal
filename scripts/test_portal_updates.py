"""ネットワーク・本番Excel・実動画を使わない更新機能の回帰テスト。"""

import contextlib
import io
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from googleapiclient.errors import HttpError
from httplib2 import Response
from openpyxl import Workbook, load_workbook

import portal_history as history
import sync
import upload_youtube as uploader


HEADERS = [
    "no.", "ファイル名", "日時", "大会名", "プレー種別",
    "選手", "ポジション", "課題", "youtube ID",
]


def make_workbook(path, count=3):
    wb = Workbook()
    ws = wb.active
    ws.append(["管理表"])
    ws.append(HEADERS)
    for number in range(1, count + 1):
        ws.append([
            number, f"{number}.mp4", datetime(2026, 9, 26), "大会A",
            "レシーブ", "選手A", "後衛", "判断", None,
        ])
    wb.save(path)
    return wb


def record(video_id, batch="one", timestamp="2026-10-03T20:00:00+09:00",
           order=1, event="大会A", event_date="2026-09-26"):
    return dict(zip(history.LOG_HEADERS, (
        batch, timestamp, video_id, timestamp, order, event_date, event,
    )))


def video(video_id, published_at):
    return {
        "id": video_id, "title": video_id, "publishedAt": published_at,
        "eventDate": "2020-01-01", "categories": {"大会": ["大会A"]},
    }


class PortalOutputTests(unittest.TestCase):
    def test_upload_dates_not_event_dates_determine_order(self):
        videos = [
            video("old", "2026-10-01T20:00:00+09:00"),
            video("new", "2026-10-02T12:00:00Z"),
            video("missing", None), video("invalid", "bad"),
            video("naive", "2026-10-03T12:00:00"),
        ]
        videos[1]["eventDate"] = "2010-01-01"
        result = sync.sort_videos(videos)
        self.assertEqual([v["id"] for v in result[:2]], ["new", "old"])
        self.assertEqual({v["id"] for v in result[2:]}, {"missing", "invalid", "naive"})
        self.assertEqual(videos[0]["id"], "old")

    def test_equal_timestamps_have_repeatable_order(self):
        a = video("a", "2026-10-03T12:00:00Z")
        b = video("b", "2026-10-03T21:00:00+09:00")
        self.assertEqual(sync.sort_videos([b, a]), sync.sort_videos([a, b]))

    def test_latest_run_only_with_multiple_events_and_no_day_merge(self):
        records = [
            record("old", batch="old", timestamp="2026-10-03T18:00:00+09:00"),
            record("a", order=1),
            record("b", order=2, event="練習試合", event_date="2026-09-23"),
            record("c", order=3, event="練習試合", event_date="2026-09-24"),
        ]
        news = history.build_latest_news(records, {"old", "a", "b", "c"})
        self.assertEqual(len(news), 1)
        self.assertEqual(news[0]["batchId"], "one")
        self.assertEqual(news[0]["videoCount"], 3)
        self.assertEqual(len(news[0]["events"]), 3)
        self.assertEqual(news, history.build_latest_news(records, {"old", "a", "b", "c"}))

    def test_incomplete_run_is_held_until_all_videos_are_visible(self):
        records = [record("a"), record("b", order=2)]
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(history.build_latest_news(records, {"a"}), [])
        self.assertEqual(history.build_latest_news(records, {"a", "b"})[0]["videoCount"], 2)

    def test_unsaved_id_excluded_and_no_history_invented_for_legacy(self):
        with contextlib.redirect_stdout(io.StringIO()):
            output = sync.make_portal_output([video("saved", None), video("unsaved", None)], {"saved"}, [])
        self.assertEqual(output["videoCount"], 1)
        self.assertEqual(output["news"], [])
        self.assertEqual(output["videos"][0]["id"], "saved")

    def test_json_replace_failure_leaves_old_file_and_cleans_temp(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "videos.json"
            original = '{"videos": []}\n'
            path.write_text(original, encoding="utf-8")
            with patch.object(sync.os, "replace", side_effect=OSError("locked")):
                with self.assertRaises(OSError):
                    sync.write_portal_output({"videos": [], "news": []}, path)
            self.assertEqual(path.read_text(encoding="utf-8"), original)
            self.assertEqual(list(Path(directory).glob(".videos-*.tmp")), [])

    def test_json_utf8_and_no_rewrite_on_identical_sync(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "videos.json"
            output = {"videos": [], "news": [{"eventName": "津山市新人戦"}]}
            sync.write_portal_output(output, path)
            self.assertEqual(json.loads(path.read_text(encoding="utf-8")), output)
            with patch.object(sync.os, "replace") as replace:
                sync.write_portal_output(output, path)
                replace.assert_not_called()


class UploadSafetyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / "videos.xlsx"
        self.wb = make_workbook(self.path)
        self.files = []
        for number in range(1, 4):
            path = self.root / f"{number}.mp4"
            path.write_bytes(b"test-only")
            self.files.append(path)
        self.stack = contextlib.ExitStack()
        self.stack.enter_context(contextlib.redirect_stdout(io.StringIO()))
        self.stack.enter_context(patch.object(uploader, "EXCEL_PATH", self.path))
        self.stack.enter_context(patch.object(uploader, "PROJECT_DIR", self.root))
        self.stack.enter_context(patch.object(uploader, "get_youtube", return_value=object()))
        self.stack.enter_context(patch.object(uploader, "find_upload_matches", side_effect=lambda name: [self.root / name]))
        self.archive = self.stack.enter_context(patch.object(uploader, "archive_video"))
        self.upload = self.stack.enter_context(patch.object(uploader, "upload_video", side_effect=["id1", "id2", "id3"]))
        self.run = self.stack.enter_context(patch.object(uploader.subprocess, "run"))
        self.run.return_value.returncode = 0
        self.input = self.stack.enter_context(patch("builtins.input", side_effect=["ALL", "N"]))

    def tearDown(self):
        self.stack.close()
        self.wb.close()
        self.temp.cleanup()

    def test_id_and_hidden_log_are_saved_before_archive(self):
        def check_saved(_):
            wb = load_workbook(self.path)
            try:
                saved, records = history.read_workbook_state(wb)
                self.assertEqual(len(saved), self.archive.call_count)
                self.assertEqual(len(saved), len(records))
                self.assertEqual(wb[history.LOG_SHEET].sheet_state, "hidden")
                self.assertEqual(wb.active.title, "Sheet")
            finally:
                wb.close()
        self.archive.side_effect = check_saved
        self.assertEqual(uploader.upload_pending(self.wb), 0)
        saved, records = history.read_excel_state(self.path)
        self.assertEqual(saved, {"id1", "id2", "id3"})
        self.assertEqual(len({r["batchId"] for r in records}), 1)
        self.assertEqual([r["orderInBatch"] for r in records], [1, 2, 3])
        self.assertEqual(len(list((self.root / "_portal_upload_backups").glob("*.xlsx"))), 1)
        self.run.assert_not_called()

    def test_excel_save_failure_stops_before_archive_and_next_upload(self):
        with patch.object(self.wb, "save", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                uploader.upload_pending(self.wb)
        self.assertEqual(self.upload.call_count, 1)
        self.archive.assert_not_called()
        self.run.assert_not_called()
        self.assertEqual(history.read_excel_state(self.path), (set(), []))

    def test_api_error_keeps_only_successful_ids_and_stops(self):
        error = HttpError(Response({"status": "403"}), b'{"error":{"message":"quota"}}')
        self.upload.side_effect = ["id1", error, "must-not-upload"]
        self.assertEqual(uploader.upload_pending(self.wb), 1)
        self.assertEqual(self.upload.call_count, 2)
        self.assertEqual(self.archive.call_count, 1)
        saved, records = history.read_excel_state(self.path)
        self.assertEqual(saved, {"id1"})
        self.assertEqual(history.build_latest_news(records, saved)[0]["videoCount"], 1)

    def test_archive_error_keeps_id_log_and_stops_remaining(self):
        self.archive.side_effect = OSError("destination exists")
        self.assertEqual(uploader.upload_pending(self.wb), 1)
        self.assertEqual(self.upload.call_count, 1)
        saved, records = history.read_excel_state(self.path)
        self.assertEqual(saved, {"id1"})
        self.assertEqual(len(records), 1)
        self.run.assert_not_called()

    def test_uploaded_rows_never_upload_again_and_pending_update_offered(self):
        self.assertEqual(uploader.upload_pending(self.wb), 0)
        self.upload.reset_mock()
        self.archive.reset_mock()
        self.input.side_effect = ["Y"]
        self.assertEqual(uploader.upload_pending(self.wb), 0)
        self.upload.assert_not_called()
        self.run.assert_called_once()
        self.assertEqual(len(history.read_excel_state(self.path)[1]), 3)

    def test_unchanged_portal_does_not_prompt_or_upload(self):
        self.assertEqual(uploader.upload_pending(self.wb), 0)
        saved, records = history.read_excel_state(self.path)
        output = sync.make_portal_output([video(v, None) for v in saved], saved, records)
        sync.write_portal_output(output, self.root / "data" / "videos.json")
        self.input.reset_mock()
        self.upload.reset_mock()
        self.assertEqual(uploader.upload_pending(self.wb), 0)
        self.input.assert_not_called()
        self.upload.assert_not_called()

    def test_cancel_and_missing_file_do_not_create_log_or_upload(self):
        self.input.side_effect = ["N"]
        self.assertEqual(uploader.upload_pending(self.wb), 0)
        self.upload.assert_not_called()
        self.assertNotIn(history.LOG_SHEET, self.wb.sheetnames)
        self.files[1].unlink()
        self.assertEqual(uploader.upload_pending(self.wb), 1)
        self.upload.assert_not_called()

    def test_invalid_log_stops_before_network(self):
        self.wb.create_sheet(history.LOG_SHEET).append(["wrong schema"])
        with self.assertRaises(ValueError):
            uploader.upload_pending(self.wb)
        self.upload.assert_not_called()


if __name__ == "__main__":
    unittest.main()
