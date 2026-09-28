from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

EXCEL_PATH = Path(
    r"C:\Users\susuk\OneDrive\勝北VBC\動画ポータル\動画管理.xlsx"
)

VIDEO_DIR = Path(
    r"C:\Users\susuk\OneDrive\勝北VBC\動画ポータル\2026\20260926_津山市新人戦\upload"
)


def split_values(value):
    """カンマ区切りのセルをリストへ変換"""
    if value is None:
        return []

    return [
        item.strip()
        for item in str(value).split(",")
        if item.strip()
    ]


def format_date(value):
    """Excelの日付を日本語表記へ変換"""
    if isinstance(value, datetime):
        return f"{value.year}年{value.month}月{value.day}日"

    # 念のため文字列の場合にも対応
    text = str(value).strip()

    try:
        dt = datetime.strptime(text, "%Y/%m/%d")
        return f"{dt.year}年{dt.month}月{dt.day}日"
    except ValueError:
        return text


def make_tags(row_data):
    """Excel 1行分からYouTubeタグを生成"""

    tags = []

    # 単一値
    if row_data["大会名"]:
        tags.append(f'大会:{row_data["大会名"]}')

    if row_data["プレー種別"]:
        tags.append(f'プレー:{row_data["プレー種別"]}')

    # 複数値
    for player in split_values(row_data["選手"]):
        tags.append(f"選手:{player}")

    for position in split_values(row_data["ポジション"]):
        tags.append(f"ポジション:{position}")

    for issue in split_values(row_data["課題"]):
        tags.append(f"課題:{issue}")

    return tags


def main():

    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"管理表が見つかりません:\n{EXCEL_PATH}"
        )

    if not VIDEO_DIR.exists():
        raise FileNotFoundError(
            f"動画フォルダが見つかりません:\n{VIDEO_DIR}"
        )

    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active

    # 2行目が見出し
    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[2]
        if cell.value is not None
    }

    required = [
        "no.",
        "ファイル名",
        "日時",
        "大会名",
        "プレー種別",
        "選手",
        "ポジション",
        "課題",
        "youtube ID",
    ]

    for name in required:
        if name not in headers:
            raise ValueError(
                f"必要な列「{name}」が見つかりません"
            )

    pending = []

    for row in range(3, ws.max_row + 1):

        filename = ws.cell(
            row, headers["ファイル名"]
        ).value

        if not filename:
            continue

        youtube_id = ws.cell(
            row, headers["youtube ID"]
        ).value

        # アップロード済み
        if youtube_id:
            continue

        no_value = ws.cell(
            row, headers["no."]
        ).value

        # 1 → 001
        no = f"{int(no_value):03d}"

        video_path = VIDEO_DIR / str(filename)

        row_data = {
            "no": no,
            "ファイル名": str(filename),
            "日時": ws.cell(
                row, headers["日時"]
            ).value,
            "大会名": ws.cell(
                row, headers["大会名"]
            ).value,
            "プレー種別": ws.cell(
                row, headers["プレー種別"]
            ).value,
            "選手": ws.cell(
                row, headers["選手"]
            ).value,
            "ポジション": ws.cell(
                row, headers["ポジション"]
            ).value,
            "課題": ws.cell(
                row, headers["課題"]
            ).value,
        }

        date_text = format_date(
            row_data["日時"]
        )

        title = (
            f'{date_text} '
            f'{row_data["大会名"]}'
            f'｜プレー振り返り {no}'
        )

        tags = make_tags(row_data)

        pending.append(
            {
                "no": no,
                "file": video_path,
                "title": title,
                "tags": tags,
            }
        )

    print()
    print("=" * 70)
    print("YouTubeアップロード予定内容")
    print("=" * 70)

    for item in pending:

        print()
        print(f'[{item["no"]}]')
        print(f'ファイル: {item["file"].name}')

        if item["file"].exists():
            print("ファイル確認: OK")
        else:
            print("ファイル確認: ★見つかりません")

        print()
        print("タイトル:")
        print(item["title"])

        print()
        print("タグ:")

        for tag in item["tags"]:
            print(f"  {tag}")

        print()
        print("-" * 70)

    print()
    print(f"アップロード予定: {len(pending)}本")


if __name__ == "__main__":
    main()