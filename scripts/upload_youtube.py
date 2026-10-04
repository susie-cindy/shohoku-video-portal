from pathlib import Path
from datetime import datetime
import os
import shutil
import subprocess
import sys
from uuid import uuid4
from openpyxl import load_workbook

from portal_history import (
    EXCEL_PATH, append_upload_record, has_pending_portal_update,
    now_jst, read_workbook_state,
)

from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload
from googleapiclient.errors import HttpError

# =========================================================
# 設定
# =========================================================

PROJECT_DIR = Path(__file__).resolve().parent.parent

CLIENT_SECRET = PROJECT_DIR / "client_secret.json"
TOKEN_FILE = PROJECT_DIR / "token.json"

VIDEO_ROOT = Path(
    r"C:\Users\susuk\OneDrive\勝北VBC\動画ポータル"
)
# アップロード・読み取り権限
GOOGLE_API = "www." + "googleapis." + "com"

SCOPES = [
    "https:" + "//" + GOOGLE_API + "/auth/youtube.upload",
    "https:" + "//" + GOOGLE_API + "/auth/youtube.readonly",
]

# =========================================================
# 補助関数
# =========================================================

def find_upload_matches(filename):
    # Excelの値をパスやglobパターンとして扱わず、ファイル名で照合する。
    if not filename or Path(filename).name != filename:
        return []
    matches = [
        folder / filename
        for folder in VIDEO_ROOT.glob("*/*/upload")
        if (folder / filename).is_file()
    ]
    if len(matches) > 1:
        raise RuntimeError(
            f"同名の動画ファイルが複数見つかりました: {filename}\n"
            + "\n".join(str(path) for path in matches)
        )
    return matches


def archive_video(video_path):
    if video_path.parent.name != "upload":
        raise ValueError(f"upload内の動画ではありません: {video_path}")
    archive = video_path.parent.parent / "archive"
    destination = archive / video_path.name
    # Windowsのrenameは移動先が存在すると失敗する（事前確認後の競合も保護）。
    # 上書き動作の異なるOSでは、安全のため移動しない。
    if os.name != "nt":
        raise OSError("archiveへの安全な移動はWindowsで実行してください。")
    archive.mkdir(exist_ok=True)
    if os.path.lexists(destination):
        raise FileExistsError(f"移動先が既に存在します: {destination}")
    video_path.rename(destination)
    print(f"archiveへ移動しました: {video_path} → {destination}")


def split_values(value):
    if value is None:
        return []

    return [
        item.strip()
        for item in str(value).split(",")
        if item.strip()
    ]


def format_date(value):
    if isinstance(value, datetime):
        return f"{value.year}年{value.month}月{value.day}日"

    text = str(value).strip()

    try:
        dt = datetime.strptime(text, "%Y/%m/%d")
        return f"{dt.year}年{dt.month}月{dt.day}日"
    except ValueError:
        return text


def make_tags(data):

    tags = []

    if data["大会名"]:
        tags.append(f'大会:{data["大会名"]}')

    if data["プレー種別"]:
        tags.append(f'プレー:{data["プレー種別"]}')

    for value in split_values(data["選手"]):
        tags.append(f"選手:{value}")

    for value in split_values(data["ポジション"]):
        tags.append(f"ポジション:{value}")

    for value in split_values(data["課題"]):
        tags.append(f"課題:{value}")

    return tags


def get_youtube():

    creds = None

    if TOKEN_FILE.exists():
        creds = Credentials.from_authorized_user_file(
            TOKEN_FILE,
            SCOPES
        )

    if not creds or not creds.valid:

        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())

        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                CLIENT_SECRET,
                SCOPES
            )

            creds = flow.run_local_server(port=0)

        TOKEN_FILE.write_text(
            creds.to_json(),
            encoding="utf-8"
        )

    return build(
        "youtube",
        "v3",
        credentials=creds
    )


# =========================================================
# メイン
# =========================================================

def upload_video(youtube, video_path, title, tags):

    body = {
        "snippet": {
            "title": title,
            "description": """勝北VBCの選手・保護者向けのプレー振り返り動画です。
通し再生 → 問題箇所で静止して確認 → もう一度通し再生、の順に見直しましょう。
本動画はチーム内での振り返りを目的として限定公開しています。
チーム外へのリンク転送、SNSへの掲載、公開再生リストへの追加はお控えください。""",
            "tags": tags,
            "categoryId": "17",   # Sports
        },
        "status": {
            "privacyStatus": "unlisted",
            "selfDeclaredMadeForKids": False,
        },
    }

    media = MediaFileUpload(
        str(video_path),
        chunksize=8 * 1024 * 1024,
        resumable=True
    )

    print("進捗 0%")

    request = youtube.videos().insert(
        part="snippet,status",
        body=body,
        media_body=media,
    )

    response = None

    while response is None:

        status, response = request.next_chunk()

        if status:
            percent = int(
                status.progress() * 100
            )
            print(
                f"\r進捗: {percent}%",
                end="",
                flush=True
            )

    video_id = response["id"]
    if not video_id:
        raise ValueError("YouTubeから有効なvideo IDが返されませんでした")
    print("\r進捗 100%", flush=True)

    print("\n\nアップロード完了")
    print(f"YouTube ID: {video_id}")

    return video_id


def main():
    wb = load_workbook(EXCEL_PATH)
    try:
        return upload_pending(wb)
    finally:
        wb.close()


def upload_pending(wb):
    # 記録形式の不整合はYouTubeへ送信する前に検出する。
    read_workbook_state(wb)
    ws = wb.active
    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[2]
        if cell.value is not None
    }
    required = [
        "no.", "ファイル名", "日時", "大会名", "プレー種別",
        "選手", "ポジション", "課題", "youtube ID",
    ]
    for name in required:
        if name not in headers:
            raise ValueError(f"必要な列「{name}」がありません")

    targets = []
    for row in range(3, ws.max_row + 1):
        # IDのある行は必ずスキップする。
        youtube_id = ws.cell(row, headers["youtube ID"]).value
        if youtube_id is not None and str(youtube_id).strip() != "":
            filename = str(ws.cell(row, headers["ファイル名"]).value or "")
            for video_path in find_upload_matches(filename):
                try:
                    archive_video(video_path)
                except OSError as exc:
                    print(f"警告: アップロード済み動画のarchive整理に失敗: {video_path}\n{exc}")
                    print("削除・上書きは行いません。この動画は再アップロードしません。")
            continue
        no_value = ws.cell(row, headers["no."]).value
        if no_value is None or str(no_value).strip() == "":
            continue
        no = f"{int(no_value):03d}"
        data = {
            name: ws.cell(row, headers[name]).value
            for name in required[1:-1]
        }
        if not str(data["大会名"] or "").strip() or not str(data["プレー種別"] or "").strip():
            raise ValueError(f"No. {no}: 大会名とプレー種別を入力してください。")
        filename = str(data["ファイル名"] or "")

        matches = find_upload_matches(filename)

        if len(matches) == 1:
            video_path = matches[0]
        elif len(matches) == 0: 
            video_path = VIDEO_ROOT / "__NOT_FOUND__" / filename

        date_text = format_date(data["日時"])
        title = (
            f'{date_text} '
            f'{data["大会名"]}'
            f'｜プレー振り返り {no}'
        )
        targets.append((row, no, filename, video_path, title, make_tags(data)))

    if not targets:
        print("未アップロードの動画はありません。")
        if has_pending_portal_update(wb, PROJECT_DIR / "data" / "videos.json"):
            print("Excel保存済みの未反映動画・履歴があります。")
            return offer_portal_update()
        return 0

    print("=" * 60)
    print("YouTube一括アップロード対象（限定公開）")
    print("No.\tファイル名\tタイトル")
    for row, no, filename, video_path, title, tags in targets:
        print(f"{no}\t{filename}\t{title}")
    print("=" * 60)

    missing = [target for target in targets if not target[3].is_file()]
    if missing:
        print("次の動画ファイルが存在しないため、アップロードを開始しません。")
        for row, no, filename, video_path, title, tags in missing:
            print(f"No. {no}: {video_path}")
        return 1

    total = len(targets)
    answer = input(
        f"{total}本をアップロードします。開始する場合は ALL と入力してください: "
    )
    if answer != "ALL":
        print("キャンセルしました。")
        return 0

    batch_id = str(uuid4())
    started_at = now_jst()
    # 初回保存や途中停止から復旧できるよう、実行前のブックを保存する。
    backup_dir = EXCEL_PATH.parent / "_portal_upload_backups"
    backup_dir.mkdir(exist_ok=True)
    backup_file = backup_dir / f"{EXCEL_PATH.stem}_{batch_id}.xlsx"
    shutil.copy2(EXCEL_PATH, backup_file)
    print(f"管理表のバックアップ: {backup_file}")
    youtube = get_youtube()
    success_count = 0
    failed_no = None
    for index, (row, no, filename, video_path, title, tags) in enumerate(targets, 1):
        print(f"\n[{index}/{total}] {no} アップロード開始")
        try:
            video_id = upload_video(youtube, video_path, title, tags)
        except HttpError as exc:
            failed_no = no
            print(f"\nNo. {no} YouTube APIエラー: {exc}")
            print("この動画のIDは書き込まず、残りのアップロードを停止します。")
            break
        except Exception as exc:
            print(f"\nNo. {no} アップロードエラー: {exc}")
            print("この動画のIDは書き込まず、処理を停止します。")
            raise

        # 次の動画へ進む前に、この動画のIDを必ず保存する。
        ws.cell(row, headers["youtube ID"]).value = video_id
        try:
            append_upload_record(
                wb, batch_id=batch_id, started_at=started_at,
                video_id=video_id, confirmed_at=now_jst(),
                order=success_count + 1,
                event_date=ws.cell(row, headers["日時"]).value,
                event_name=ws.cell(row, headers["大会名"]).value,
            )
            wb.save(EXCEL_PATH)
        except Exception as exc:
            print(f"\nNo. {no} Excel保存エラー: {exc}")
            print(f"アップロード済みのYouTube ID: {video_id}")
            print("処理を停止します。再実行前に、このIDをExcelへ登録してください。")
            raise
        success_count += 1
        print("動画管理.xlsx にYouTube IDを保存しました。")
        print(f"https://youtu.be/{video_id}")
        try:
            archive_video(video_path)
        except OSError as exc:
            print(f"\nNo. {no} archive移動エラー: {exc}")
            print("YouTubeアップロードとYouTube IDのExcel保存は完了済みです。再アップロードしてはいけません。")
            print(f"YouTube ID: {video_id} / 元動画: {video_path}")
            print(f"成功（ID保存済み）：{success_count}本")
            print("archive移動失敗：1本。残りの処理を停止します。")
            print(f"未実行：{total - index}本")
            return 1

    print(f"\n成功（ID保存済み）：{success_count}本")
    if failed_no is not None:
        print(f"失敗：1本（No. {failed_no}）")
        print(f"未実行：{total - success_count - 1}本")
    else:
        print(f"{total}本すべてのアップロードとID保存が完了しました。")

    if success_count > 0:
        if offer_portal_update() != 0:
            return 1
    return 1 if failed_no is not None else 0


def offer_portal_update():
    answer = input("ポータルを更新しますか？ [Y/N] ")
    if answer not in ("Y", "y"):
        print("ポータル更新を見送りました。")
        return 0
    try:
        result = subprocess.run(
            [sys.executable, str(PROJECT_DIR / "scripts" / "update_portal.py")],
            cwd=PROJECT_DIR,
        )
    except OSError as exc:
        print(f"ポータル更新の起動に失敗しました: {exc}")
        return 1
    if result.returncode != 0:
        print("ポータル更新に失敗しました。保存済みのYouTube IDは保持されています。")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
